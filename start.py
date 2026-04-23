import os
import glob
from datetime import datetime  # <- [NUEVO] Importamos el modulo de tiempo
from openpyxl import load_workbook, Workbook
from openpyxl.styles import PatternFill

# === CONFIGURACION GENERAL ===
TIENDAS = ['CrazyFamily', 'OfertasImper', 'Abizi', 'Moisess']
CARPETA_SALIDA = '.- Archivos maestros'

def limpiar_texto(v):
    if v is None or str(v).strip() == "" or str(v).lower() == "nan":
        return ""
    t = str(v).strip()
    return t[:-2] if t.endswith('.0') else t

def limpiar_precio(v):
    if v is None or str(v).strip() == "" or str(v).lower() == "nan":
        return ""
    t = str(v).strip()
    if t.endswith('.0'):
        t = t[:-2]
    try:
        return int(t)
    except ValueError:
        try:
            return float(t.replace(',', '.'))
        except ValueError:
            return t

def consolidar_listados():
    print("--- [*] INICIANDO CONSOLIDACION (CON HISTORIAL POR FECHA) ---")
    ruta_base = os.getcwd()
    ruta_hechos = os.path.join(ruta_base, CARPETA_SALIDA)

    if not os.path.exists(ruta_hechos):
        os.makedirs(ruta_hechos)

    # --- [NUEVO] GENERACION DE NOMBRE DINAMICO ---
    # Formato: AñoMesDia_HoraMinutoSegundo (Ej: 20260422_153025)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    nombre_archivo = f'LISTADO_NORMALIZADO_{timestamp}.xlsx'
    ruta_archivo_final = os.path.join(ruta_hechos, nombre_archivo)

    # Como siempre será un nombre nuevo, creamos el Excel desde cero
    print(f"[*] Creando nuevo archivo maestro: {nombre_archivo}")
    wb_out = Workbook()
    if 'Sheet' in wb_out.sheetnames:
        wb_out.remove(wb_out['Sheet'])
    # ---------------------------------------------

    rojo = PatternFill(start_color='FF0000', end_color='FF0000', fill_type='solid')
    total_registros_procesados = 0

    for tienda in TIENDAS:
        ruta_tienda = os.path.join(ruta_base, tienda)
        print(f"\n[-] Revisando zona: {tienda}...")

        if not os.path.exists(ruta_tienda):
            os.makedirs(ruta_tienda)
            continue

        archivos_excel = glob.glob(os.path.join(ruta_tienda, '*.xlsx'))
        archivos_origen = [f for f in archivos_excel if not os.path.basename(f).startswith('~$')]

        if not archivos_origen:
            print(f"  [>] Sin archivos nuevos en '{tienda}'.")
            continue

        archivo_in = archivos_origen[0]
        print(f"  [v] Extrayendo de: {os.path.basename(archivo_in)}")

        try:
            wb_ml = load_workbook(archivo_in, data_only=True)
            ws_ml = wb_ml['Publicaciones'] if 'Publicaciones' in wb_ml.sheetnames else wb_ml.active

            fila_encabezados = 5 
            for r in range(1, 10): 
                for c in range(1, ws_ml.max_column + 1):
                    val = str(ws_ml.cell(row=r, column=c).value).strip().lower()
                    if val == "sku":
                        fila_encabezados = r
                        break
                if fila_encabezados == r:
                    break
            
            print(f"  [i] Encabezados detectados en la fila {fila_encabezados}.")

            c_sku, c_tit, c_pre, c_est = None, None, None, None
            
            for col in range(1, ws_ml.max_column + 1):
                raw_header = ws_ml.cell(row=fila_encabezados, column=col).value
                if raw_header is None:
                    continue
                    
                header = str(raw_header).replace('\xa0', ' ').strip().lower()
                
                if c_sku is None and header == "sku":
                    c_sku = col
                elif c_tit is None and header in ["título", "titulo"]:
                    c_tit = col
                elif c_pre is None and header == "precio":
                    c_pre = col
                elif c_est is None and header in ["estado", "estado de publicación", "status", "condición"]:
                    c_est = col

            if c_est is None:
                print(f"  [!] ALERTA EN {tienda}: No encontre 'Estado' en la fila {fila_encabezados}.")
            
            c_sku = c_sku or 5
            c_tit = c_tit or 6
            c_pre = c_pre or 10
            c_est = c_est or 23
            
            print(f"  [=] Columnas Ancladas -> SKU:{c_sku}, Titulo:{c_tit}, Precio:{c_pre}, Estado:{c_est}")

            productos = []
            titulo_memoria = None
            precio_memoria = None
            estado_memoria = None

            for r in range(6, ws_ml.max_row + 1):
                raw_sku = ws_ml.cell(row=r, column=c_sku).value
                raw_tit = ws_ml.cell(row=r, column=c_tit).value
                raw_pre = ws_ml.cell(row=r, column=c_pre).value
                raw_est = ws_ml.cell(row=r, column=c_est).value

                if all(v is None or str(v).strip() == "" for v in [raw_sku, raw_tit, raw_pre, raw_est]):
                    titulo_memoria = None
                    precio_memoria = None
                    estado_memoria = None
                    continue

                if raw_tit is not None and str(raw_tit).strip() != "":
                    titulo_memoria = raw_tit
                if raw_pre is not None and str(raw_pre).strip() != "":
                    precio_memoria = raw_pre
                if raw_est is not None and str(raw_est).strip() != "":
                    estado_memoria = raw_est

                tit_final = raw_tit if (raw_tit is not None and str(raw_tit).strip() != "") else titulo_memoria
                pre_final = raw_pre if (raw_pre is not None and str(raw_pre).strip() != "") else precio_memoria
                est_final = raw_est if (raw_est is not None and str(raw_est).strip() != "") else estado_memoria

                fila_procesada = [
                    limpiar_texto(raw_sku), 
                    limpiar_texto(tit_final), 
                    limpiar_precio(pre_final), 
                    limpiar_texto(est_final)
                ]

                if raw_sku is not None or raw_tit is not None:
                    productos.append(fila_procesada)

            # Inyeccion en el Maestro
            if tienda in wb_out.sheetnames:
                ws_out = wb_out[tienda]
            else:
                ws_out = wb_out.create_sheet(title=tienda)
                ws_out.append(["SKU", "Título", "Precio", "Estado"])

            for i, p_data in enumerate(productos):
                row_idx = i + 2
                for col_idx, valor in enumerate(p_data, start=1):
                    celda = ws_out.cell(row=row_idx, column=col_idx, value=valor)
                    if valor == "":
                        celda.fill = rojo

            print(f"  [+] {len(productos)} productos inyectados en la pestana '{tienda}'.")
            total_registros_procesados += len(productos)

        except Exception as e:
            print(f"  [x] ERROR al procesar '{tienda}': {e}")

    print("\n[*] Guardando archivo maestro...")
    try:
        wb_out.save(ruta_archivo_final)
        print(f"[OK] EXITO! Se consolidaron {total_registros_procesados} productos en total.")
    except PermissionError:
        print(f"[x] ERROR: No se pudo guardar. Asegurate de tener cerrado el archivo resultante.")

if __name__ == "__main__":
    consolidar_listados()